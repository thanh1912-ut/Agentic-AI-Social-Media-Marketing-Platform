"""API tests for tenant-scoped campaigns, immutable versions and approval gates."""

from __future__ import annotations

import asyncio
import csv
import io
from datetime import datetime, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from database.models import Base, CampaignPost, PostVersion, new_id
from services.api import campaign_workflows
from services.api.db import get_db
from services.api.main import app
from services.api.storage import LocalObjectStorage


@pytest.fixture
def workflow_api():
    engine = create_async_engine(
        "sqlite+aiosqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    session_factory = async_sessionmaker(engine, expire_on_commit=False)

    async def setup():
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)

    async def override_db():
        async with session_factory() as session:
            yield session

    asyncio.run(setup())
    app.dependency_overrides[get_db] = override_db
    try:
        with TestClient(app) as client:
            yield client, session_factory
    finally:
        app.dependency_overrides.clear()
        asyncio.run(engine.dispose())


def _register(client: TestClient, email: str) -> dict:
    response = client.post("/api/v1/auth/register", json={
        "email": email,
        "password": "safe-test-password",
        "full_name": "Test Owner",
        "company_name": f"Company {email}",
    })
    assert response.status_code == 201, response.text
    return response.json()


def test_campaign_api_has_tenant_scoping_and_versioned_approval(workflow_api) -> None:
    client, _session_factory = workflow_api
    owner = _register(client, "owner-campaign@example.com")
    workspace_id = owner["active_workspace_id"]
    created = client.post(f"/api/v1/workspaces/{workspace_id}/campaigns", headers={"X-CSRF-Token": client.cookies.get("agentic_csrf")}, json={
        "name": "Bếp Mộc mùa hè",
        "brief": {
            "objective": "engagement",
            "audience": ["Gia đình địa phương"],
            "product_ids": [],
            "key_message": "Món Việt dễ chọn cho bữa cơm gia đình",
            "must_include": [],
            "must_avoid": ["cam kết chữa bệnh"],
            "start_date": "2026-09-01",
            "end_date": "2026-09-30",
        },
        "pillars": ["product"],
        "channels": ["facebook_page"],
    })
    assert created.status_code == 201, created.text
    campaign = created.json()
    assert campaign["post_count"] == 0

    other = TestClient(app)
    with other:
        second = _register(other, "other-campaign@example.com")
        denied = other.get(f"/api/v1/workspaces/{workspace_id}/campaigns/{campaign['id']}")
        assert denied.status_code == 404
        assert denied.json()["error"]["code"] == "not_found"

    csrf = client.cookies.get("agentic_csrf")
    headers = {"X-CSRF-Token": csrf}
    created_post = client.post(
        f"/api/v1/workspaces/{workspace_id}/campaigns/{campaign['id']}/posts",
        headers=headers,
        json={"pillar": "product", "format": "text", "caption": "Món ngon cho cả nhà", "hashtags": ["#BepMoc"]},
    )
    assert created_post.status_code == 201, created_post.text
    post_id = created_post.json()["id"]
    post_url = f"/api/v1/workspaces/{workspace_id}/posts/{post_id}"

    first_edit = client.patch(post_url, headers=headers, json={"version": 1, "caption": "Bài đã được biên tập"})
    assert first_edit.status_code == 200, first_edit.text
    assert first_edit.json()["version"] == 2
    assert first_edit.json()["current"]["caption"] == "Bài đã được biên tập"

    stale = client.patch(post_url, headers=headers, json={"version": 1, "caption": "Bản cũ"})
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "version_conflict"

    submitted = client.post(f"{post_url}/submit-approval", headers=headers, json={"version": 2})
    assert submitted.status_code == 200, submitted.text
    assert submitted.json()["status"] == "needs_review"
    approved = client.post(f"{post_url}/approval", headers=headers, json={"version": 2, "decision": "approved"})
    assert approved.status_code == 200, approved.text
    assert approved.json()["current"]["approved_by"] == owner["user"]["id"]

    second_edit = client.patch(post_url, headers=headers, json={"version": 2, "caption": "Bài chỉnh sau khi duyệt"})
    assert second_edit.status_code == 200, second_edit.text
    assert second_edit.json()["version"] == 3
    assert second_edit.json()["requires_reapproval"] is True
    versions = client.get(f"{post_url}/versions").json()
    assert versions["current_version"] == 3
    assert versions["versions"][0]["caption"] == "Bài chỉnh sau khi duyệt"
    assert versions["versions"][-1]["caption"] == "Món ngon cho cả nhà"


def test_content_generation_fails_closed_until_provider_data_flow_is_approved(workflow_api) -> None:
    client, _session_factory = workflow_api
    owner = _register(client, "content-blocker@example.com")
    workspace_id = owner["active_workspace_id"]
    response = client.post(
        f"/api/v1/workspaces/{workspace_id}/posts/generate",
        headers={"X-CSRF-Token": client.cookies.get("agentic_csrf")},
        json={"campaign_id": new_id(), "count": 1},
    )
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "provider_approval_required"


def test_export_writes_downloadable_csv_and_is_idempotent(workflow_api, tmp_path, monkeypatch) -> None:
    client, session_factory = workflow_api
    owner = _register(client, "export-owner@example.com")
    workspace_id = owner["active_workspace_id"]
    csrf = client.cookies.get("agentic_csrf")
    headers = {"X-CSRF-Token": csrf}
    campaign_response = client.post(f"/api/v1/workspaces/{workspace_id}/campaigns", headers=headers, json={
        "name": "Bếp Mộc",
        "brief": {
            "objective": "engagement",
            "audience": ["Gia đình"],
            "product_ids": [],
            "key_message": "Món nhà",
            "must_include": [],
            "must_avoid": [],
            "start_date": "2026-09-01",
            "end_date": "2026-09-30",
        },
        "pillars": ["product"],
        "channels": ["facebook_page"],
    })
    assert campaign_response.status_code == 201, campaign_response.text
    campaign_id = campaign_response.json()["id"]
    post_id = new_id()
    now = datetime.now(timezone.utc)
    content = {"caption": "=1+1, món nhà", "hashtags": ["#BepMoc"], "media": [], "citations": []}

    async def seed():
        async with session_factory() as db:
            db.add(CampaignPost(
                id=post_id, company_id=workspace_id, campaign_id=campaign_id,
                channel="facebook_page", pillar="product", format="text", status="draft",
                current_version=1, current_json=content, created_at=now, updated_at=now,
            ))
            db.add(PostVersion(
                id=new_id(), company_id=workspace_id, campaign_id=campaign_id,
                post_id=post_id, version=1, content_json=content, source="human",
                created_by=owner["user"]["id"], created_by_name="Test Owner", created_at=now,
            ))
            await db.commit()

    asyncio.run(seed())
    monkeypatch.setattr(campaign_workflows, "storage", LocalObjectStorage(tmp_path))
    export_headers = {**headers, "Idempotency-Key": "export-test-idempotency-001"}
    body = {"campaign_id": campaign_id, "format": "csv"}
    created = client.post(f"/api/v1/workspaces/{workspace_id}/exports", headers=export_headers, json=body)
    assert created.status_code == 202, created.text
    job = created.json()["job"]
    assert job["status"] == "succeeded"
    export_id = job["result"]["export_id"]

    repeated = client.post(f"/api/v1/workspaces/{workspace_id}/exports", headers=export_headers, json=body)
    assert repeated.status_code == 202
    assert repeated.json()["job_id"] == created.json()["job_id"]

    metadata = client.get(f"/api/v1/workspaces/{workspace_id}/exports/{export_id}")
    assert metadata.status_code == 200
    assert metadata.json()["size"] > 0
    download = client.get(f"/api/v1/workspaces/{workspace_id}/exports/{export_id}/download")
    assert download.status_code == 200
    assert download.headers["content-disposition"].startswith("attachment;")
    rows = list(csv.reader(io.StringIO(download.content.decode("utf-8-sig"))))
    assert rows[0][6] == "Nội dung"
    assert rows[1][6].startswith("'=1+1")

    xlsx_response = client.post(
        f"/api/v1/workspaces/{workspace_id}/exports",
        headers={**headers, "Idempotency-Key": "export-xlsx-idempotency-001"},
        json={"campaign_id": campaign_id, "format": "xlsx"},
    )
    assert xlsx_response.status_code == 202, xlsx_response.text
    xlsx_id = xlsx_response.json()["job"]["result"]["export_id"]
    xlsx_download = client.get(f"/api/v1/workspaces/{workspace_id}/exports/{xlsx_id}/download")
    assert xlsx_download.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_download.content), read_only=True, data_only=True)
    assert workbook.active.cell(row=2, column=7).value.startswith("'=1+1")
