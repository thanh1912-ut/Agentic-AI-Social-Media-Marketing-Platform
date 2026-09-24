"""Parser contracts for the supported document formats."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from services.ingestion.parsers import ParseError, parse_document


def test_docx_preserves_paragraph_and_table_locators(tmp_path: Path) -> None:
    path = tmp_path / "brand.docx"
    document = DocxDocument()
    document.add_paragraph("Bếp Mộc phục vụ món Việt.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Sản phẩm"
    table.cell(0, 1).text = "Giá"
    table.cell(1, 0).text = "Cơm gà"
    table.cell(1, 1).text = "65.000đ"
    document.save(path)

    parsed = parse_document(
        path,
        kind="docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=path.name,
    )

    assert [(block.text, block.locator) for block in parsed.text_blocks] == [
        ("Bếp Mộc phục vụ món Việt.", "paragraph=1")
    ]
    assert len(parsed.table_blocks) == 1
    assert parsed.table_blocks[0].headers == ["Sản phẩm", "Giá"]
    assert parsed.table_blocks[0].rows == [["Cơm gà", "65.000đ"]]
    assert parsed.table_blocks[0].locator == "table=1;row=1"
    assert parsed.metadata == {"paragraphs": 1, "tables": 1}


def test_docx_with_only_a_table_is_valid_content(tmp_path: Path) -> None:
    path = tmp_path / "menu.docx"
    document = DocxDocument()
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Món"
    table.cell(1, 0).text = "Bún chả"
    document.save(path)

    parsed = parse_document(path, kind="docx", mime_type="application/octet-stream", filename=path.name)

    assert not parsed.text_blocks
    assert parsed.table_blocks[0].rows == [["Bún chả"]]


def test_xlsx_returns_each_sheet_as_a_located_table(tmp_path: Path) -> None:
    path = tmp_path / "menu.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Món Việt"
    sheet.append(["Sản phẩm", "Giá"])
    sheet.append(["Cơm gà", "65.000đ"])
    second = workbook.create_sheet("Đồ uống")
    second.append(["Tên", "Dung tích"])
    second.append(["Cà phê sữa", "350ml"])
    workbook.save(path)

    parsed = parse_document(
        path,
        kind="xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=path.name,
    )

    assert [table.locator for table in parsed.table_blocks] == [
        "sheet=Món Việt;row=1",
        "sheet=Đồ uống;row=1",
    ]
    assert parsed.table_blocks[0].rows == [["Cơm gà", "65.000đ"]]
    assert parsed.table_blocks[1].rows == [["Cà phê sữa", "350ml"]]
    assert parsed.metadata == {"rows": 4, "sheets": 2}


def test_csv_and_txt_preserve_vietnamese_content_and_locators(tmp_path: Path) -> None:
    csv_path = tmp_path / "menu.csv"
    csv_path.write_text("Sản phẩm,Giá\nCơm gà,65.000đ\n", encoding="utf-8")
    csv_parsed = parse_document(csv_path, kind="csv", mime_type="text/csv", filename=csv_path.name)
    assert csv_parsed.table_blocks[0].headers == ["Sản phẩm", "Giá"]
    assert csv_parsed.table_blocks[0].rows == [["Cơm gà", "65.000đ"]]
    assert csv_parsed.table_blocks[0].locator == "csv:row=1"

    txt_path = tmp_path / "brand.txt"
    txt_path.write_text("Bếp Mộc, giọng nói thân thiện.", encoding="utf-8")
    txt_parsed = parse_document(txt_path, kind="txt", mime_type="text/plain", filename=txt_path.name)
    assert txt_parsed.text_blocks[0].text == "Bếp Mộc, giọng nói thân thiện."
    assert txt_parsed.text_blocks[0].locator == "text:1"


def test_csv_skips_empty_rows_and_rejects_files_without_any_cells(tmp_path: Path) -> None:
    path = tmp_path / "blank-rows.csv"
    path.write_text("\n, ,\nProduct,Price\nCoffee,5\n", encoding="utf-8")
    parsed = parse_document(path, kind="csv", mime_type="text/csv", filename=path.name)
    assert parsed.table_blocks[0].headers == ["Product", "Price"]
    assert parsed.table_blocks[0].rows == [["Coffee", "5"]]

    empty_path = tmp_path / "empty.csv"
    empty_path.write_text(", ,\n ,\n", encoding="utf-8")
    with pytest.raises(ParseError) as error:
        parse_document(empty_path, kind="csv", mime_type="text/csv", filename=empty_path.name)
    assert error.value.code == "empty_content"


def test_parser_reports_corrupt_empty_and_unsupported_inputs(tmp_path: Path) -> None:
    empty = tmp_path / "empty.txt"
    empty.write_text("  ", encoding="utf-8")
    with pytest.raises(ParseError) as empty_error:
        parse_document(empty, kind="txt", mime_type="text/plain", filename=empty.name)
    assert empty_error.value.code == "empty_content"

    corrupt_xlsx = tmp_path / "corrupt.xlsx"
    corrupt_xlsx.write_bytes(b"not a zip file")
    with pytest.raises(ParseError) as corrupt_error:
        parse_document(corrupt_xlsx, kind="xlsx", mime_type="application/octet-stream", filename=corrupt_xlsx.name)
    assert corrupt_error.value.code == "corrupted"

    unknown = tmp_path / "notes.rtf"
    unknown.write_text("unsupported", encoding="utf-8")
    with pytest.raises(ParseError) as unsupported_error:
        parse_document(unknown, kind="rtf", mime_type="application/rtf", filename=unknown.name)
    assert unsupported_error.value.code == "unsupported_type"


def _write_text_pdf(path: Path) -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 50 700 Td (Brand facts) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as output:
        writer.write(output)


def test_pdf_text_has_page_locator(tmp_path: Path) -> None:
    path = tmp_path / "brand.pdf"
    _write_text_pdf(path)

    parsed = parse_document(path, kind="pdf", mime_type="application/pdf", filename=path.name)

    assert parsed.text_blocks[0].text == "Brand facts"
    assert parsed.text_blocks[0].locator == "page=1"
    assert parsed.metadata == {"pages": 1, "characters": 11}


def test_csv_and_pdf_enforce_parser_limits(tmp_path: Path, monkeypatch) -> None:
    from services.ingestion import parsers

    csv_path = tmp_path / "large.csv"
    csv_path.write_text("header\nrow1\nrow2\n", encoding="utf-8")
    monkeypatch.setattr(parsers, "MAX_TABLE_ROWS", 1)
    with pytest.raises(ParseError) as row_error:
        parse_document(csv_path, kind="csv", mime_type="text/csv", filename=csv_path.name)
    assert row_error.value.code == "parser_limit_exceeded"

    pdf_path = tmp_path / "many-pages.pdf"
    _write_text_pdf(pdf_path)
    monkeypatch.setattr(parsers, "MAX_PDF_PAGES", 0)
    with pytest.raises(ParseError) as page_error:
        parse_document(pdf_path, kind="pdf", mime_type="application/pdf", filename=pdf_path.name)
    assert page_error.value.code == "parser_limit_exceeded"


def test_xlsx_and_docx_enforce_uncompressed_archive_limit(tmp_path: Path, monkeypatch) -> None:
    from services.ingestion import parsers

    xlsx_path = tmp_path / "limited.xlsx"
    workbook = Workbook()
    workbook.active.append(["Header"])
    workbook.active.append(["Value"])
    workbook.save(xlsx_path)
    docx_path = tmp_path / "limited.docx"
    docx = DocxDocument()
    docx.add_paragraph("Content")
    docx.save(docx_path)

    monkeypatch.setattr(parsers, "MAX_ARCHIVE_UNCOMPRESSED_BYTES", 1)
    for path, kind, mime_type in (
        (xlsx_path, "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        (docx_path, "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ):
        with pytest.raises(ParseError) as error:
            parse_document(path, kind=kind, mime_type=mime_type, filename=path.name)
        assert error.value.code == "parser_limit_exceeded"


def test_txt_csv_xlsx_and_image_enforce_output_limits(tmp_path: Path, monkeypatch) -> None:
    from PIL import Image

    from services.ingestion import parsers

    txt_path = tmp_path / "large.txt"
    txt_path.write_text("12345", encoding="utf-8")
    monkeypatch.setattr(parsers, "MAX_PARSED_TEXT_CHARACTERS", 4)
    with pytest.raises(ParseError) as text_error:
        parse_document(txt_path, kind="txt", mime_type="text/plain", filename=txt_path.name)
    assert text_error.value.code == "parser_limit_exceeded"

    csv_path = tmp_path / "wide.csv"
    csv_path.write_text("first,second\na,b\n", encoding="utf-8")
    monkeypatch.setattr(parsers, "MAX_PARSED_TEXT_CHARACTERS", 100)
    monkeypatch.setattr(parsers, "MAX_TABLE_COLUMNS", 1)
    with pytest.raises(ParseError) as csv_error:
        parse_document(csv_path, kind="csv", mime_type="text/csv", filename=csv_path.name)
    assert csv_error.value.code == "parser_limit_exceeded"

    xlsx_path = tmp_path / "many-cells.xlsx"
    workbook = Workbook()
    workbook.active.append(["Header"])
    workbook.active.append(["Value"])
    workbook.save(xlsx_path)
    monkeypatch.setattr(parsers, "MAX_TABLE_COLUMNS", 256)
    monkeypatch.setattr(parsers, "MAX_TABLE_CELLS", 1)
    with pytest.raises(ParseError) as xlsx_error:
        parse_document(
            xlsx_path,
            kind="xlsx",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=xlsx_path.name,
        )
    assert xlsx_error.value.code == "parser_limit_exceeded"

    image_path = tmp_path / "large.png"
    Image.new("RGB", (2, 2), color="white").save(image_path)
    monkeypatch.setattr(parsers, "MAX_IMAGE_PIXELS", 3)
    with pytest.raises(ParseError) as image_error:
        parse_document(image_path, kind="image", mime_type="image/png", filename=image_path.name)
    assert image_error.value.code == "parser_limit_exceeded"
