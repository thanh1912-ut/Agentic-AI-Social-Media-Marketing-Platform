"""Isolated Docling conversion process for untrusted uploaded documents."""
from __future__ import annotations

import csv
import importlib.metadata
import io
import json
import os
import sys
import tempfile
import traceback
import zipfile
from pathlib import Path
from typing import Any

from services.ingestion.parsers import (
    MAX_ARCHIVE_ENTRIES,
    MAX_ARCHIVE_UNCOMPRESSED_BYTES,
    MAX_PARSED_TEXT_CHARACTERS,
    MAX_PDF_PAGES,
    MAX_TABLE_CELLS,
    MAX_TABLE_COLUMNS,
    MAX_TABLE_ROWS,
    ParseError,
    ParsedDocument,
    ParsedTableBlock,
    ParsedTextBlock,
    is_image_signature,
)


CSV_DELIMITERS = ",;\t|"
def _limits(argument: str) -> dict[str, int]:
    values = json.loads(argument)
    defaults = {
        "text": MAX_PARSED_TEXT_CHARACTERS,
        "rows": MAX_TABLE_ROWS,
        "columns": MAX_TABLE_COLUMNS,
        "cells": MAX_TABLE_CELLS,
        "pages": MAX_PDF_PAGES,
        "archive_bytes": MAX_ARCHIVE_UNCOMPRESSED_BYTES,
        "archive_entries": MAX_ARCHIVE_ENTRIES,
    }
    if not isinstance(values, dict):
        raise ValueError("invalid parser limits")
    for name, value in values.items():
        if name not in defaults or not isinstance(value, int) or value < 0:
            raise ValueError("invalid parser limits")
        defaults[name] = value
    return defaults


def _fail(code: str, filename: str, hint: str, message: str | None = None) -> ParseError:
    return ParseError(code, message or f"Không đọc được tệp “{filename}”.", hint)


def _check_zip(path: Path, filename: str, limits: dict[str, int]) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
    except Exception as exc:
        raise _fail("corrupted", filename, "Hãy mở và lưu lại tệp rồi tải lên lần nữa.") from exc
    if len(entries) > limits["archive_entries"]:
        raise _fail("parser_limit_exceeded", filename, f"Tệp có tối đa {limits['archive_entries']} thành phần.")
    if sum(item.file_size for item in entries) > limits["archive_bytes"]:
        raise _fail("parser_limit_exceeded", filename, "Kích thước tệp sau giải nén vượt giới hạn an toàn.")


def _kind_signature(path: Path, kind: str, filename: str, limits: dict[str, int]) -> None:
    with path.open("rb") as source:
        head = source.read(16)
    if is_image_signature(head):
        raise _fail("unsupported_type", filename, "Mục Tài liệu chỉ nhận PDF có lớp chữ, DOCX, XLSX, CSV và TXT.")
    if kind == "pdf" and not head.startswith(b"%PDF-"):
        raise _fail("type_mismatch", filename, "Phần mở rộng không khớp nội dung file.")
    if kind in {"docx", "xlsx"}:
        _check_zip(path, filename, limits)
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
        except Exception as exc:
            raise _fail("corrupted", filename, "Hãy mở và lưu lại tệp Office rồi tải lên lần nữa.") from exc
        marker = "word/document.xml" if kind == "docx" else "xl/workbook.xml"
        if "[Content_Types].xml" not in names or marker not in names:
            raise _fail("type_mismatch", filename, "Phần mở rộng không khớp nội dung file.")
    elif kind in {"csv", "txt"} and b"\x00" in head:
        raise _fail("type_mismatch", filename, "Tệp có dấu hiệu là dữ liệu nhị phân, không phải văn bản.")
    if kind == "image":
        raise _fail("unsupported_type", filename, "Ảnh không được hỗ trợ trong mục Tài liệu.")


def _require_pdf_models(filename: str) -> None:
    artifact_value = os.environ.get("DOCLING_ARTIFACTS_PATH", "").strip()
    if not artifact_value:
        raise _fail("parser_model_unavailable", filename, "Quản trị viên cần cài trước model Docling layout/table và đặt DOCLING_ARTIFACTS_PATH.")
    root = Path(artifact_value).resolve()
    manifest_path = root / "docling-models.sha256.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        files = manifest["files"]
        if (
            manifest.get("docling_version") != importlib.metadata.version("docling-slim")
            or not {"layout", "tableformer"}.issubset(set(manifest.get("model_families", [])))
            or not files
        ):
            raise ValueError("incompatible model manifest")
        for relative, checksum in files.items():
            file_path = (root / relative).resolve()
            if root not in file_path.parents or not file_path.is_file() or len(checksum) != 64:
                raise ValueError("invalid model file entry")
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError, importlib.metadata.PackageNotFoundError) as exc:
        raise _fail("parser_model_unavailable", filename, "Model Docling local thiếu hoặc sai phiên bản. Hãy chạy bước verify model rồi khởi động lại worker.") from exc


def _csv_encoding(path: Path, filename: str) -> str:
    sample = path.read_bytes()[:65536]
    for encoding in ("utf-8-sig", "utf-8", "cp1258"):
        try:
            sample.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    try:
        sample.decode("latin-1")
        return "latin-1"
    except UnicodeDecodeError as exc:
        raise _fail("corrupted", filename, "Hãy xuất CSV ở định dạng UTF-8.") from exc


def _csv_dialect(path: Path, encoding: str):
    with path.open("r", encoding=encoding, newline="") as source:
        sample = source.read(65536)
    try:
        return csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS)
    except csv.Error:
        return csv.excel


def _table_from_item(item: Any, locator: str) -> ParsedTableBlock | None:
    data = item.data
    if data.num_rows <= 0 or data.num_cols <= 0:
        return None
    grid = data.grid
    matrix: list[list[str]] = []
    for row_index, row in enumerate(grid):
        values = []
        for column_index, cell in enumerate(row):
            # Preserve a merged cell's value once. Repeating it into each
            # covered cell would fabricate values in a table.
            covered = (
                row_index > cell.start_row_offset_idx
                or column_index > cell.start_col_offset_idx
            )
            values.append("" if covered and (cell.row_span > 1 or cell.col_span > 1) else str(cell.text or "").strip())
        matrix.append(values)
    if not matrix:
        return None
    headers = [value or f"column_{index + 1}" for index, value in enumerate(matrix[0])]
    return ParsedTableBlock(headers=headers, rows=matrix[1:], locator=locator)


def _convert_one(converter, path: Path, kind: str, filename: str, limits: dict[str, int]) -> ParsedDocument:
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.document import ConversionStatus
    from docling_core.types.doc import TableItem, TextItem

    from docling.datamodel.document import DocumentStream

    stream = DocumentStream(name=filename, stream=io.BytesIO(path.read_bytes()))
    result = converter.convert(stream, max_file_size=25 * 1024 * 1024, max_num_pages=limits["pages"])
    if result.status == ConversionStatus.FAILURE:
        detail = "; ".join(error.error_message for error in result.errors[:3])
        raise _fail("corrupted", filename, "Kiểm tra file có bị hỏng, khóa mật khẩu hoặc sai định dạng không.", detail or None)
    texts: list[ParsedTextBlock] = []
    tables: list[ParsedTableBlock] = []
    warnings: list[str] = []
    characters = 0
    page_numbers: set[int] = set()
    for item, _level in result.document.iterate_items():
        if isinstance(item, TableItem):
            prov = item.prov or []
            page = next((entry.page_no for entry in prov if entry.page_no is not None), None)
            locator = f"page={page};table={len(tables) + 1}" if page is not None else f"table={len(tables) + 1}"
            table = _table_from_item(item, locator)
            if table is not None:
                tables.append(table)
                characters += sum(map(len, table.headers)) + sum(len(cell) for row in table.rows for cell in row)
                page_numbers.update(entry.page_no for entry in prov if entry.page_no is not None)
        elif isinstance(item, TextItem):
            text = item.text.strip()
            if not text:
                continue
            prov = item.prov or []
            page = next((entry.page_no for entry in prov if entry.page_no is not None), None)
            if page is not None:
                page_numbers.add(page)
            label = getattr(item.label, "value", str(item.label))
            locator = f"page={page};item={len(texts) + 1}" if page is not None else f"item={len(texts) + 1}"
            heading = text if label in {"section_header", "title"} else None
            texts.append(ParsedTextBlock(text=text, locator=locator, heading=heading))
            characters += len(text)
    if characters > limits["text"]:
        raise _fail("parser_limit_exceeded", filename, f"Văn bản chuẩn hóa tối đa {limits['text']:,} ký tự.")
    if result.status == ConversionStatus.PARTIAL_SUCCESS:
        warnings.append("docling_partial_conversion")
    if result.errors:
        warnings.extend("docling_conversion_warning" for _ in result.errors[:8])
    if kind == "pdf":
        if not texts and not tables:
            raise _fail("pdf_no_text_layer", filename, "PDF scan không có lớp chữ. OCR đang tắt; hãy tải PDF có thể chọn/copy chữ.")
        total_pages = len(result.document.pages)
        if total_pages > limits["pages"]:
            raise _fail("parser_limit_exceeded", filename, f"PDF được nhận tối đa {limits['pages']} trang.")
        if total_pages and page_numbers and len(page_numbers) < total_pages:
            warnings.append("pdf_pages_without_extractable_text")
    if not texts and not tables:
        raise _fail("empty_content", filename, "Hãy chọn tài liệu có văn bản hoặc bảng.")
    row_count = sum(len(table.rows) for table in tables)
    cell_count = sum(len(table.headers) + sum(len(row) for row in table.rows) for table in tables)
    if row_count > limits["rows"]:
        raise _fail("parser_limit_exceeded", filename, f"Tài liệu được nhận tối đa {limits['rows']:,} dòng dữ liệu.")
    if any(len(table.headers) > limits["columns"] for table in tables):
        raise _fail("parser_limit_exceeded", filename, f"Mỗi bảng được nhận tối đa {limits['columns']} cột.")
    if cell_count > limits["cells"]:
        raise _fail("parser_limit_exceeded", filename, f"Tài liệu được nhận tối đa {limits['cells']:,} ô bảng.")
    return ParsedDocument(
        text_blocks=texts,
        table_blocks=tables,
        metadata={
            "engine": "docling",
            "engine_version": "2.130.0",
            "characters": characters,
            "text_blocks": len(texts),
            "tables": len(tables),
            "rows": row_count,
            "columns": max((len(table.headers) for table in tables), default=0),
            "cells": cell_count,
            "pages": len(result.document.pages) if kind == "pdf" else None,
            "pictures_ignored": True,
        },
        warnings=warnings,
    )


def _convert_csv(path: Path, filename: str, converter, limits: dict[str, int]) -> ParsedDocument:
    encoding = _csv_encoding(path, filename)
    dialect = _csv_dialect(path, encoding)
    total_rows = 0
    total_cells = 0
    total_characters = 0
    headers: list[str] | None = None
    with path.open("r", encoding=encoding, newline="") as source:
        reader = csv.reader(source, dialect)
        for row_number, row in enumerate(reader, start=1):
            if not row or not any(cell.strip() for cell in row):
                continue
            if headers is None:
                headers = [cell.strip() or f"column_{i + 1}" for i, cell in enumerate(row)]
                width = len(headers)
                if width > limits["columns"]:
                    raise _fail("parser_limit_exceeded", filename, f"Bảng được nhận tối đa {limits['columns']} cột.")
            else:
                if len(row) != width:
                    raise _fail("csv_row_width_mismatch", filename, f"Dòng dữ liệu {row_number} có số cột khác header; hãy kiểm tra dấu phân cách và dấu ngoặc kép.")
                total_rows += 1
                total_cells += len(row)
                total_characters += sum(len(cell) for cell in row)
            if total_rows > limits["rows"] or total_cells + width > limits["cells"]:
                raise _fail("parser_limit_exceeded", filename, f"CSV vượt giới hạn {limits['rows']:,} dòng hoặc {limits['cells']:,} ô.")
            if total_characters > limits["text"]:
                raise _fail("parser_limit_exceeded", filename, f"Văn bản CSV tối đa {limits['text']:,} ký tự.")
    if headers is None:
        raise _fail("empty_content", filename, "CSV không có header hoặc dòng dữ liệu.")

    data_rows: list[list[str]] = []
    tables: list[ParsedTableBlock] = []
    source_record = 0
    batch_start_record = 0
    saw_header = False
    with path.open("r", encoding=encoding, newline="") as source:
        reader = csv.reader(source, dialect)
        for record_number, row in enumerate(reader, start=1):
            source_record = record_number
            if not saw_header:
                if row and any(cell.strip() for cell in row):
                    saw_header = True
                continue
            if not row or not any(cell.strip() for cell in row):
                continue
            if not batch_start_record:
                batch_start_record = record_number
            data_rows.append(row)
            cells_per_batch = max(1, (20_000 // max(width, 1)) - 1)
            if len(data_rows) >= min(1_000, cells_per_batch):
                _convert_csv_batch(data_rows, headers, batch_start_record, record_number, filename, converter, tables)
                data_rows.clear()
                batch_start_record = 0
        if data_rows:
            _convert_csv_batch(data_rows, headers, batch_start_record, source_record, filename, converter, tables)

    # Preserve a header-only CSV as a one-row, zero-data table.
    if not tables:
        tables.append(ParsedTableBlock(headers=headers, rows=[], locator="csv:row=1"))
    emitted_rows = sum(len(table.rows) for table in tables)
    emitted_cells = sum(sum(len(row) for row in table.rows) for table in tables)
    if emitted_rows != total_rows or emitted_cells != total_cells or any(len(table.headers) != width for table in tables):
        raise _fail("docling_table_integrity_failed", filename, "Docling không trả đủ dòng/ô CSV đã đọc trước khi chia lô.")
    return ParsedDocument(
        table_blocks=tables,
        metadata={"engine": "docling", "engine_version": "2.130.0", "rows": total_rows, "columns": width, "cells": total_cells + width, "tables": len(tables), "delimiter": dialect.delimiter, "characters": total_characters + sum(map(len, headers)), "pictures_ignored": True},
    )


def _convert_csv_batch(rows, headers, start_source_row, final_source_row, filename, converter, tables):
    with tempfile.TemporaryDirectory(prefix="docling-csv-") as directory:
        batch_path = Path(directory) / "part.csv"
        with batch_path.open("w", encoding="utf-8", newline="") as out:
            writer = csv.writer(out)
            writer.writerow(headers)
            writer.writerows(rows)
        result = converter.convert(batch_path, max_file_size=25 * 1024 * 1024)
        from docling_core.types.doc import TableItem

        emitted_rows = 0
        for item, _level in result.document.iterate_items():
            if isinstance(item, TableItem):
                matrix = _table_from_item(item, f"csv:row={start_source_row}-{final_source_row}")
                if matrix is None:
                    continue
                # Docling recognizes the repeated batch header. Use the checked
                # original header to keep columns stable across every chunk.
                tables.append(ParsedTableBlock(headers=headers, rows=matrix.rows, locator=matrix.locator, warnings=matrix.warnings))
                emitted_rows += len(matrix.rows)
        if emitted_rows != len(rows):
            raise _fail(
                "docling_table_integrity_failed",
                filename,
                "Docling không trả đủ số dòng cho một lô CSV; dữ liệu chưa được chấp nhận.",
            )


def _convert_xlsx(path: Path, filename: str, converter, limits: dict[str, int]) -> ParsedDocument:
    _check_zip(path, filename, limits)
    source_file = path.open("rb")
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(source_file, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        source_file.close()
        raise _fail("corrupted", filename, "Hãy mở và lưu lại tệp XLSX rồi tải lên lần nữa.") from exc
    data_only = None
    data_file = None
    try:
        data_file = path.open("rb")
        data_only = openpyxl.load_workbook(data_file, read_only=True, data_only=True, keep_links=False)
        tables: list[ParsedTableBlock] = []
        warnings: list[str] = []
        total_rows = total_cells = total_data_cells = total_characters = 0
        total_rows_emitted = 0
        total_cells_emitted = 0
        for sheet_index, sheet in enumerate(workbook.worksheets):
            if sheet.max_column and sheet.max_column > limits["columns"]:
                raise _fail("parser_limit_exceeded", filename, f"Mỗi bảng được nhận tối đa {limits['columns']} cột.")
            values_sheet = data_only.worksheets[sheet_index]
            source_rows = sheet.iter_rows(values_only=True)
            values_rows = values_sheet.iter_rows(values_only=True)
            rows_buffer: list[list[Any]] = []
            headers: list[str] | None = None
            source_start = 1
            sheet_data_rows = 0
            for row_number, (formula_row, value_row) in enumerate(zip(source_rows, values_rows, strict=False), start=1):
                normalized: list[str] = []
                for formula, value in zip(formula_row, value_row, strict=False):
                    if isinstance(formula, str) and formula.startswith("="):
                        if value is None:
                            normalized.append(f"[Công thức chưa tính: {formula}]")
                            warnings.append("xlsx_formula_value_missing")
                        else:
                            normalized.append(str(value))
                    elif value is None:
                        normalized.append("")
                    else:
                        normalized.append(str(value))
                if not any(cell.strip() for cell in normalized):
                    continue
                if headers is None:
                    headers = [cell or f"column_{i + 1}" for i, cell in enumerate(normalized)]
                    total_cells += len(headers)
                    total_characters += sum(map(len, headers))
                    if total_cells > limits["cells"] or total_characters > limits["text"]:
                        raise _fail("parser_limit_exceeded", filename, "Bảng tính vượt giới hạn số ô hoặc số ký tự đã công bố.")
                    source_start = row_number + 1
                    rows_buffer = []
                    continue
                total_rows += 1
                sheet_data_rows += 1
                total_cells += len(normalized)
                total_data_cells += len(normalized)
                total_characters += sum(map(len, normalized))
                if total_rows > limits["rows"] or total_cells > limits["cells"] or total_characters > limits["text"]:
                    raise _fail("parser_limit_exceeded", filename, "Bảng tính vượt giới hạn số dòng, số ô hoặc số ký tự đã công bố.")
                rows_buffer.append(normalized)
                batch_cells = (len(rows_buffer) + 1) * len(headers)
                if len(rows_buffer) >= min(1_000, max(1, 20_000 // len(headers) - 1)):
                    emitted_rows, emitted_cells = _convert_xlsx_batch(rows_buffer, headers, sheet.title, source_start, row_number, filename, converter, tables)
                    total_rows_emitted += emitted_rows
                    total_cells_emitted += emitted_cells
                    rows_buffer.clear()
                    source_start = row_number + 1
            if headers is not None and rows_buffer:
                emitted_rows, emitted_cells = _convert_xlsx_batch(rows_buffer, headers, sheet.title, source_start, row_number, filename, converter, tables)
                total_rows_emitted += emitted_rows
                total_cells_emitted += emitted_cells
            if headers is not None and sheet_data_rows == 0:
                tables.append(ParsedTableBlock(headers=headers, rows=[], locator=f"sheet={sheet.title};row=1"))
        workbook.close()
        data_only.close()
        data_file.close()
        source_file.close()
    except ParseError:
        workbook.close()
        if data_only is not None:
            data_only.close()
        if data_file is not None:
            data_file.close()
        source_file.close()
        raise
    except Exception as exc:
        workbook.close()
        if data_only is not None:
            data_only.close()
        if data_file is not None:
            data_file.close()
        source_file.close()
        raise _fail("corrupted", filename, "Hãy kiểm tra cấu trúc workbook, công thức hoặc sheet rồi lưu lại XLSX.") from exc
    if not tables:
        raise _fail("empty_content", filename, "XLSX không có ô dữ liệu.")
    if total_rows_emitted != total_rows or total_cells_emitted != total_data_cells:
        raise _fail("docling_table_integrity_failed", filename, "Docling không trả đủ ô/dòng XLSX đã đọc trước khi chia lô.")
    return ParsedDocument(table_blocks=tables, metadata={"engine": "docling", "engine_version": "2.130.0", "rows": total_rows, "columns": max((len(table.headers) for table in tables), default=0), "cells": total_cells, "characters": total_characters, "sheets": len(workbook.sheetnames), "tables": len(tables), "pictures_ignored": True}, warnings=list(dict.fromkeys(warnings)))


def _convert_xlsx_batch(rows, headers, sheet_name, original_start, original_end, filename, converter, tables):
    import openpyxl
    from docling_core.types.doc import TableItem

    with tempfile.TemporaryDirectory(prefix="docling-xlsx-") as directory:
        batch_path = Path(directory) / "part.xlsx"
        output = openpyxl.Workbook(write_only=True)
        sheet = output.create_sheet(title=sheet_name[:31] or "Sheet1")
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        output.save(batch_path)
        result = converter.convert(batch_path, max_file_size=25 * 1024 * 1024)
        emitted_rows = emitted_cells = 0
        for item, _level in result.document.iterate_items():
            if isinstance(item, TableItem):
                parsed = _table_from_item(item, f"sheet={sheet_name};rows={original_start}-{original_end}")
                if parsed is not None:
                    parsed = ParsedTableBlock(headers=headers, rows=parsed.rows, locator=parsed.locator, warnings=parsed.warnings)
                    tables.append(parsed)
                    emitted_rows += len(parsed.rows)
                    emitted_cells += sum(len(row) for row in parsed.rows)
        if emitted_rows != len(rows) or emitted_cells != len(rows) * len(headers):
            raise _fail("docling_table_integrity_failed", filename, f"Docling không giữ đủ dòng/ô của sheet {sheet_name}.")
        return emitted_rows, emitted_cells


def _convert_document(path: Path, kind: str, filename: str, limits: dict[str, int]) -> ParsedDocument:
    if kind not in {"pdf", "docx", "xlsx", "csv", "txt"}:
        raise _fail("unsupported_type", filename, "Chỉ hỗ trợ PDF có lớp chữ, DOCX, XLSX, CSV và TXT. Ảnh không được hỗ trợ.")
    _kind_signature(path, kind, filename, limits)
    if kind == "pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            if reader.is_encrypted:
                raise _fail("encrypted", filename, "Hãy gỡ mật khẩu PDF rồi tải lại.")
            if len(reader.pages) > limits["pages"]:
                raise _fail("parser_limit_exceeded", filename, f"PDF được nhận tối đa {limits['pages']} trang.")
            if not any((page.extract_text() or "").strip() for page in reader.pages):
                raise _fail("pdf_no_text_layer", filename, "PDF scan không có lớp chữ. OCR đang tắt; hãy tải PDF có thể chọn/copy chữ.")
        except ParseError:
            raise
        except Exception as exc:
            raise _fail("corrupted", filename, "Hãy xuất lại PDF rồi tải tệp lên lần nữa.") from exc
        _require_pdf_models(filename)
    if kind == "csv":
        # One converter instance is reused for all bounded CSV batches.
        converter = _converter()
        return _convert_csv(path, filename, converter, limits)
    if kind == "xlsx":
        converter = _converter()
        return _convert_xlsx(path, filename, converter, limits)
    if kind == "pdf":
        converter = _converter()
    else:
        converter = _converter()
    result = _convert_one(converter, path, kind, filename, limits)
    if kind == "txt":
        # Text files can use the Markdown backend; preserve punctuation,
        # headings and code-like content through the document item's text.
        if result.metadata.get("characters", 0) == 0:
            raise _fail("empty_content", filename, "Tệp TXT không có nội dung văn bản.")
    return result


def _converter():
    from docling.datamodel.accelerator_options import AcceleratorDevice, AcceleratorOptions
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, PdfFormatOption

    artifact_path = os.environ.get("DOCLING_ARTIFACTS_PATH") or None
    pdf_options = PdfPipelineOptions(
        do_ocr=False,
        do_table_structure=True,
        do_picture_classification=False,
        do_picture_description=False,
        do_chart_extraction=False,
        do_code_enrichment=False,
        do_formula_enrichment=False,
        generate_page_images=False,
        generate_picture_images=False,
        generate_table_images=False,
        enable_remote_services=False,
        allow_external_plugins=False,
        accelerator_options=AcceleratorOptions(num_threads=4, device=AcceleratorDevice.CPU),
        artifacts_path=artifact_path,
        document_timeout=600,
    )
    return DocumentConverter(
        allowed_formats=[InputFormat.PDF, InputFormat.DOCX, InputFormat.XLSX, InputFormat.CSV, InputFormat.MD],
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_options)},
    )


def main() -> int:
    if len(sys.argv) != 6:
        return 64
    source, kind, filename, output, limits_arg = sys.argv[1:]
    source_path = Path(source)
    output_path = Path(output)
    try:
        parsed = _convert_document(source_path, kind, filename, _limits(limits_arg))
        output_path.write_text(json.dumps({"ok": True, "document": {
            "text_blocks": [block.__dict__ for block in parsed.text_blocks],
            "table_blocks": [block.__dict__ for block in parsed.table_blocks],
            "metadata": parsed.metadata,
            "warnings": parsed.warnings,
        }}, ensure_ascii=False), encoding="utf-8")
        return 0
    except ParseError as exc:
        output_path.write_text(json.dumps({"ok": False, "error": {"code": exc.code, "message": exc.message, "hint": exc.hint, "retryable": exc.retryable}}, ensure_ascii=False), encoding="utf-8")
        return 1
    except Exception as exc:
        # Diagnostic detail stays on the child pipe. The parent captures it
        # with a strict byte bound and never writes it into application logs.
        traceback.print_exc(file=sys.stderr)
        output_path.write_text(json.dumps({"ok": False, "error": {"code": "parser_model_unavailable" if kind == "pdf" else "parser_unavailable", "message": f"Không thể khởi tạo bộ Docling để đọc tệp “{filename}”.", "hint": "Kiểm tra dependency Docling và model local của worker.", "retryable": True}}, ensure_ascii=False), encoding="utf-8")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
