"""Bounded, deterministic parsers used by the background worker.

The parser never executes document content. It returns locators and warnings
for M3; authorization remains in the API/database layer.
"""

from __future__ import annotations

import csv
import io
import mimetypes
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


class ParseError(Exception):
    def __init__(self, code: str, message: str, hint: str, *, retryable: bool = False):
        super().__init__(message)
        self.code = code
        self.message = message
        self.hint = hint
        self.retryable = retryable


@dataclass(frozen=True)
class ParsedTextBlock:
    text: str
    locator: str
    heading: str | None = None
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class ParsedTableBlock:
    headers: list[str]
    rows: list[list[str]]
    locator: str
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class ParsedDocument:
    text_blocks: list[ParsedTextBlock] = field(default_factory=list)
    table_blocks: list[ParsedTableBlock] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


MAX_PARSED_TEXT_CHARACTERS = 2_000_000
MAX_TABLE_ROWS = 20_000
MAX_TABLE_COLUMNS = 256
MAX_TABLE_CELLS = 200_000
MAX_PDF_PAGES = 1_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 10_000
MAX_IMAGE_PIXELS = 40_000_000


SUPPORTED_MIME_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/csv": "csv",
    "text/plain": "txt",
    "image/jpeg": "image",
    "image/png": "image",
    "image/webp": "image",
}
SUPPORTED_EXTENSIONS = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".csv": "csv",
    ".txt": "txt",
    ".jpg": "image",
    ".jpeg": "image",
    ".png": "image",
    ".webp": "image",
}


def infer_kind(filename: str, mime_type: str | None = None) -> str | None:
    if mime_type in SUPPORTED_MIME_TYPES:
        return SUPPORTED_MIME_TYPES[mime_type]
    return SUPPORTED_EXTENSIONS.get(Path(filename).suffix.lower())


def _limit_exceeded(filename: str, detail: str) -> ParseError:
    return ParseError(
        "parser_limit_exceeded",
        f"Tệp “{filename}” vượt giới hạn xử lý an toàn.",
        detail,
    )


def _check_archive_limits(path: Path, filename: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
    except Exception as exc:
        raise ParseError("corrupted", f"Không đọc được tệp “{filename}”.", "Hãy mở và lưu lại tệp rồi tải lên lần nữa.") from exc
    if len(entries) > MAX_ARCHIVE_ENTRIES:
        raise _limit_exceeded(filename, f"Tệp lưu trữ có tối đa {MAX_ARCHIVE_ENTRIES} thành phần.")
    if sum(entry.file_size for entry in entries) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        raise _limit_exceeded(filename, "Kích thước nội dung sau giải nén vượt giới hạn xử lý.")


def _check_table_limits(filename: str, *, rows: int, columns: int, cells: int) -> None:
    if rows > MAX_TABLE_ROWS:
        raise _limit_exceeded(filename, f"Bảng được nhận tối đa {MAX_TABLE_ROWS} dòng dữ liệu.")
    if columns > MAX_TABLE_COLUMNS:
        raise _limit_exceeded(filename, f"Bảng được nhận tối đa {MAX_TABLE_COLUMNS} cột.")
    if cells > MAX_TABLE_CELLS:
        raise _limit_exceeded(filename, f"Tài liệu được nhận tối đa {MAX_TABLE_CELLS} ô bảng.")


def _decode_text(raw: bytes, filename: str) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("corrupted", f"Không đọc được tệp “{filename}”.", "Hãy xuất lại tệp ở định dạng UTF-8.")


def _parse_txt(path: Path, filename: str) -> ParsedDocument:
    text = _decode_text(path.read_bytes(), filename).strip()
    if len(text) > MAX_PARSED_TEXT_CHARACTERS:
        raise _limit_exceeded(filename, f"Văn bản được nhận tối đa {MAX_PARSED_TEXT_CHARACTERS} ký tự.")
    if not text:
        raise ParseError("empty_content", f"Tệp “{filename}” không có nội dung.", "Hãy chọn tệp có văn bản.")
    return ParsedDocument(text_blocks=[ParsedTextBlock(text=text, locator="text:1")], metadata={"characters": len(text)})


def _parse_csv(path: Path, filename: str) -> ParsedDocument:
    text = _decode_text(path.read_bytes(), filename)
    if len(text) > MAX_PARSED_TEXT_CHARACTERS:
        raise _limit_exceeded(filename, f"CSV được nhận tối đa {MAX_PARSED_TEXT_CHARACTERS} ký tự.")
    reader = csv.reader(io.StringIO(text))
    first_row = next((row for row in reader if any(cell.strip() for cell in row)), None)
    if first_row is None:
        raise ParseError("empty_content", f"Tệp “{filename}” không có dữ liệu.", "Hãy chọn tệp CSV có ít nhất một dòng dữ liệu.")
    headers = [cell.strip() or f"column_{i + 1}" for i, cell in enumerate(first_row)]
    if len(headers) > MAX_TABLE_COLUMNS:
        raise _limit_exceeded(filename, f"Bảng được nhận tối đa {MAX_TABLE_COLUMNS} cột.")
    normalized_rows: list[list[str]] = []
    cells = len(headers)
    for row in reader:
        if not any(cell.strip() for cell in row):
            continue
        if len(normalized_rows) >= MAX_TABLE_ROWS:
            raise _limit_exceeded(filename, f"Bảng được nhận tối đa {MAX_TABLE_ROWS} dòng dữ liệu.")
        if len(row) > MAX_TABLE_COLUMNS:
            raise _limit_exceeded(filename, f"Bảng được nhận tối đa {MAX_TABLE_COLUMNS} cột.")
        if len(row) > len(headers):
            added_columns = len(row) - len(headers)
            for index in range(len(headers), len(row)):
                headers.append(f"column_{index + 1}")
            for existing in normalized_rows:
                existing.extend([""] * added_columns)
            cells += added_columns * (len(normalized_rows) + 1)
        row = row + [""] * (len(headers) - len(row))
        row = row[: len(headers)]
        cells += len(row)
        _check_table_limits(filename, rows=len(normalized_rows) + 1, columns=len(headers), cells=cells)
        normalized_rows.append(row)
    return ParsedDocument(
        table_blocks=[ParsedTableBlock(headers=headers, rows=normalized_rows, locator="csv:row=1")],
        metadata={"rows": len(normalized_rows) + 1, "columns": len(headers)},
    )


def _parse_xlsx(path: Path, filename: str) -> ParsedDocument:
    try:
        import openpyxl
    except ImportError as exc:
        raise ParseError("parser_unavailable", "Chưa cài bộ đọc XLSX.", "Liên hệ quản trị viên để bật parser XLSX.", retryable=True) from exc
    _check_archive_limits(path, filename)
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError("corrupted", f"Không đọc được tệp “{filename}”.", "Hãy mở và lưu lại tệp XLSX rồi tải lên lần nữa.") from exc
    tables: list[ParsedTableBlock] = []
    total_rows = 0
    total_cells = 0
    total_characters = 0
    try:
        for sheet in workbook.worksheets:
            declared_rows = sheet.max_row or 0
            declared_columns = sheet.max_column or 0
            if declared_columns > MAX_TABLE_COLUMNS or declared_rows * declared_columns > MAX_TABLE_CELLS:
                raise _limit_exceeded(filename, "Kích thước bảng tính vượt giới hạn số dòng/cột/ô.")
            values: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                normalized = ["" if value is None else str(value) for value in row]
                if not any(cell.strip() for cell in normalized):
                    continue
                total_rows += 1
                total_cells += len(normalized)
                total_characters += sum(map(len, normalized))
                if total_characters > MAX_PARSED_TEXT_CHARACTERS:
                    raise _limit_exceeded(filename, f"Nội dung được nhận tối đa {MAX_PARSED_TEXT_CHARACTERS} ký tự.")
                _check_table_limits(filename, rows=total_rows, columns=max(map(len, normalized)), cells=total_cells)
                values.append(normalized)
            if not values:
                continue
            width = max(len(row) for row in values)
            values = [row + [""] * (width - len(row)) for row in values]
            tables.append(ParsedTableBlock(values[0], values[1:], f"sheet={sheet.title};row=1"))
    finally:
        workbook.close()
    if not tables:
        raise ParseError("empty_content", f"Tệp “{filename}” không có dữ liệu.", "Hãy chọn bảng tính có dữ liệu.")
    return ParsedDocument(table_blocks=tables, metadata={"rows": total_rows, "sheets": len(tables)})


def _parse_docx(path: Path, filename: str) -> ParsedDocument:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise ParseError("parser_unavailable", "Chưa cài bộ đọc DOCX.", "Liên hệ quản trị viên để bật parser DOCX.", retryable=True) from exc
    _check_archive_limits(path, filename)
    try:
        document = DocxDocument(path)
    except Exception as exc:
        raise ParseError("corrupted", f"Không đọc được tệp “{filename}”.", "Hãy mở và lưu lại tệp DOCX rồi tải lên lần nữa.") from exc
    blocks: list[ParsedTextBlock] = []
    total_characters = 0
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        total_characters += len(text)
        if total_characters > MAX_PARSED_TEXT_CHARACTERS:
            raise _limit_exceeded(filename, f"Văn bản được nhận tối đa {MAX_PARSED_TEXT_CHARACTERS} ký tự.")
        blocks.append(ParsedTextBlock(text=text, locator=f"paragraph={index}"))
    tables: list[ParsedTableBlock] = []
    total_rows = 0
    total_cells = 0
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            total_characters += sum(map(len, values))
            if total_characters > MAX_PARSED_TEXT_CHARACTERS:
                raise _limit_exceeded(filename, f"Văn bản được nhận tối đa {MAX_PARSED_TEXT_CHARACTERS} ký tự.")
            if any(values):
                rows.append(values)
        total_rows += len(rows)
        total_cells += sum(len(row) for row in rows)
        _check_table_limits(filename, rows=total_rows, columns=max((len(row) for row in rows), default=0), cells=total_cells)
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        headers = [cell or f"column_{index + 1}" for index, cell in enumerate(rows[0])]
        tables.append(
            ParsedTableBlock(
                headers=headers,
                rows=rows[1:],
                locator=f"table={table_index};row=1",
            )
        )
    if not blocks and not tables:
        raise ParseError("empty_content", f"Tệp “{filename}” không có văn bản hoặc bảng.", "Hãy chọn DOCX có nội dung văn bản hoặc bảng.")
    return ParsedDocument(
        text_blocks=blocks,
        table_blocks=tables,
        metadata={"paragraphs": len(blocks), "tables": len(tables)},
    )


def _parse_pdf(path: Path, filename: str) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ParseError("parser_unavailable", "Chưa cài bộ đọc PDF.", "Liên hệ quản trị viên để bật parser PDF.", retryable=True) from exc
    try:
        reader = PdfReader(str(path))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise _limit_exceeded(filename, f"PDF được nhận tối đa {MAX_PDF_PAGES} trang.")
        if reader.is_encrypted:
            raise ParseError("encrypted", f"PDF “{filename}” có mật khẩu.", "Hãy gỡ mật khẩu rồi tải tệp lên lại.")
        blocks = []
        total_characters = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            total_characters += len(text)
            if total_characters > MAX_PARSED_TEXT_CHARACTERS:
                raise _limit_exceeded(filename, f"PDF được nhận tối đa {MAX_PARSED_TEXT_CHARACTERS} ký tự văn bản.")
            if text:
                blocks.append(ParsedTextBlock(text=text, locator=f"page={page_number}"))
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("corrupted", f"Không đọc được PDF “{filename}”.", "Hãy xuất lại PDF rồi tải lên lần nữa.") from exc
    if not blocks:
        raise ParseError("pdf_no_text_layer", f"PDF “{filename}” không có lớp văn bản.", "Hãy tải bản PDF có thể chọn/copy chữ hoặc xử lý OCR trước.")
    return ParsedDocument(text_blocks=blocks, metadata={"pages": len(reader.pages), "characters": sum(len(b.text) for b in blocks)})


def _parse_image(path: Path, filename: str) -> ParsedDocument:
    try:
        from PIL import Image
        with Image.open(path) as image:
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise _limit_exceeded(filename, f"Ảnh được nhận tối đa {MAX_IMAGE_PIXELS} pixel.")
            metadata = {"width": image.width, "height": image.height, "format": image.format, "images": 1}
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("corrupted", f"Không đọc được ảnh “{filename}”.", "Hãy tải lên ảnh PNG, JPG hoặc WEBP hợp lệ.") from exc
    return ParsedDocument(metadata=metadata, warnings=["image_not_text_extracted"])


def parse_document(path: Path, *, kind: str, mime_type: str, filename: str) -> ParsedDocument:
    parser = {
        "txt": _parse_txt,
        "csv": _parse_csv,
        "xlsx": _parse_xlsx,
        "docx": _parse_docx,
        "pdf": _parse_pdf,
        "image": _parse_image,
    }.get(kind)
    if parser is None:
        guessed = infer_kind(filename, mime_type) or mimetypes.guess_type(filename)[0] or "unknown"
        raise ParseError("unsupported_type", f"Định dạng “{guessed}” chưa được hỗ trợ.", "Chỉ nhận PDF, DOCX, XLSX, CSV, TXT và ảnh.")
    return parser(path, filename)

