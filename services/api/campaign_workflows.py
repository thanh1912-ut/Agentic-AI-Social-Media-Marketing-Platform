"""Tenant-scoped campaign, generated content, versioning and approval APIs."""

from __future__ import annotations

import csv
import hashlib
import io
import json
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Depends, Header, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from database.models import (
    AuditEvent,
    Campaign,
    CampaignPost,
    ExportArtifact,
    Job,
    Membership,
    PostApproval,
    PostVersion,
    User,
    new_id,
    utcnow,
)
from .campaign_schemas import (
    ApprovalRecordOut,
    ApprovalRequest,
    CampaignCreateRequest,
    CampaignOut,
    CreateManualPostRequest,
    CreateExportRequest,
    ExportOut,
    GenerateContentRequest,
    GenerateContentResponse,
    PaginatedCampaigns,
    PaginatedPosts,
    PostOut,
    PostVersionListOut,
    PostVersionOut,
    SubmitApprovalRequest,
    UpdatePostRequest,
)
from .db import get_db
from .dependencies import current_user, membership_for, require_csrf, require_permission
from .errors import ApiProblem
from .job_service import accepted_response
from .permissions import has_permission
from .storage import S3ObjectStorage, storage


router = APIRouter(tags=["campaigns-content-approvals"])
ALLOWED_PILLARS = {
    "education", "entertainment", "inspiration", "promotion", "community",
    "behind_the_scenes", "product", "testimonial",
}
EXPORT_COLUMNS = {
    "post_id": "Mã bài",
    "version": "Phiên bản",
    "status": "Trạng thái",
    "channel": "Kênh",
    "pillar": "Trụ nội dung",
    "format": "Định dạng",
    "caption": "Nội dung",
    "hashtags": "Hashtag",
    "scheduled_at": "Lịch đăng",
    "citations": "Nguồn tham khảo",
}


def _spreadsheet_safe(value: Any) -> str:
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + text
    return text


def _export_bytes(rows: list[dict[str, Any]], columns: list[str], file_format: str) -> bytes:
    if file_format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow([EXPORT_COLUMNS[column] for column in columns])
        for row in rows:
            writer.writerow([_spreadsheet_safe(row.get(column)) for column in columns])
        return ("\ufeff" + output.getvalue()).encode("utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lịch nội dung"
    sheet.append([EXPORT_COLUMNS[column] for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
    for row in rows:
        sheet.append([_spreadsheet_safe(row.get(column)) for column in columns])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.columns:
        values = [len(str(cell.value or "")) for cell in cells]
        sheet.column_dimensions[cells[0].column_letter].width = min(max(max(values, default=10) + 2, 12), 70)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


async def _campaign_out(db: AsyncSession, row: Campaign) -> CampaignOut:
    count = int(await db.scalar(select(func.count(CampaignPost.id)).where(
        CampaignPost.company_id == row.company_id,
        CampaignPost.campaign_id == row.id,
    )) or 0)
    approved = int(await db.scalar(select(func.count(CampaignPost.id)).where(
        CampaignPost.company_id == row.company_id,
        CampaignPost.campaign_id == row.id,
        CampaignPost.status == "approved",
    )) or 0)
    published = int(await db.scalar(select(func.count(CampaignPost.id)).where(
        CampaignPost.company_id == row.company_id,
        CampaignPost.campaign_id == row.id,
        CampaignPost.status == "published",
    )) or 0)
    return CampaignOut(
        id=row.id,
        workspace_id=row.company_id,
        name=row.name,
        status=row.status,
        brief=row.brief_json,
        pillars=row.pillars_json,
        channels=row.channels_json,
        version=row.version,
        post_count=count,
        approved_count=approved,
        published_count=published,
        created_by=row.created_by,
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


async def _post_out(db: AsyncSession, row: CampaignPost) -> PostOut:
    version = await db.scalar(select(PostVersion).where(
        PostVersion.company_id == row.company_id,
        PostVersion.post_id == row.id,
        PostVersion.version == row.current_version,
    ))
    if version is None:
        raise ApiProblem(500, "content_version_missing", "Không tìm thấy phiên bản nội dung hiện tại.")
    current: dict[str, Any] = {
        **version.content_json,
        "version": version.version,
        "source": version.source,
        "created_by": version.created_by,
        "created_by_name": version.created_by_name,
        "created_at": version.created_at,
        "note": version.note,
    }
    approval = await db.scalar(select(PostApproval).where(
        PostApproval.company_id == row.company_id,
        PostApproval.post_id == row.id,
        PostApproval.version == row.current_version,
        PostApproval.decision == "approved",
    ).order_by(PostApproval.decided_at.desc()))
    if approval:
        current["approved_at"] = approval.decided_at
        current["approved_by"] = approval.decided_by
    return PostOut(
        id=row.id,
        campaign_id=row.campaign_id,
        workspace_id=row.company_id,
        channel=row.channel,
        pillar=row.pillar,
        format=row.format,
        status=row.status,
        version=row.current_version,
        current=current,
        scheduled_at=row.scheduled_at,
        publish_mode=row.publish_mode,
        requires_reapproval=row.requires_reapproval,
        rejection_reason=row.rejection_reason,
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


def _check_post_version(row: CampaignPost, supplied: int) -> None:
    if row.current_version != supplied:
        raise ApiProblem(
            409,
            "version_conflict",
            "Bài viết vừa được cập nhật. Tải lại phiên bản mới nhất trước khi sửa.",
            details={"current_version": row.current_version, "your_version": supplied},
        )


async def _get_post(db: AsyncSession, company_id: str, post_id: str, *, lock: bool = False) -> CampaignPost:
    query = select(CampaignPost).where(
        CampaignPost.company_id == company_id,
        CampaignPost.id == post_id,
    )
    if lock:
        query = query.with_for_update()
    row = await db.scalar(query)
    if row is None:
        raise ApiProblem(404, "not_found", "Không tìm thấy bài viết.")
    return row


@router.get("/workspaces/{company_id}/campaigns", response_model=PaginatedCampaigns)
async def list_campaigns(
    company_id: str,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    status: str | None = None,
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    statement = select(Campaign).where(Campaign.company_id == company_id)
    if status:
        statement = statement.where(Campaign.status == status)
    total = int(await db.scalar(select(func.count()).select_from(statement.subquery())) or 0)
    rows = (await db.scalars(statement.order_by(Campaign.updated_at.desc()).offset((page - 1) * page_size).limit(page_size))).all()
    return PaginatedCampaigns(items=[await _campaign_out(db, row) for row in rows], total=total, page=page, page_size=page_size)


@router.post(
    "/workspaces/{company_id}/campaigns",
    response_model=CampaignOut,
    status_code=201,
    dependencies=[Depends(require_csrf)],
)
async def create_campaign(
    company_id: str,
    request: CampaignCreateRequest,
    user: User = Depends(current_user),
    membership: Membership = Depends(require_permission("campaign:create")),
    db: AsyncSession = Depends(get_db),
):
    if not all(pillar in ALLOWED_PILLARS for pillar in request.pillars):
        raise ApiProblem(422, "validation_error", "Chiến dịch có trụ nội dung không được hỗ trợ.")
    now = utcnow()
    brief = request.brief.model_dump(mode="json")
    row = Campaign(
        id=new_id(),
        company_id=company_id,
        name=request.name,
        status="draft",
        brief_json=brief,
        pillars_json=request.pillars,
        channels_json=request.channels,
        version=1,
        created_by=user.id,
        created_at=now,
        updated_at=now,
    )
    db.add(row)
    db.add(AuditEvent(company_id=company_id, actor_user_id=user.id, action="campaign.create", entity_type="campaign", entity_id=row.id, metadata_json={"version": 1}))
    await db.commit()
    return await _campaign_out(db, row)


@router.get("/workspaces/{company_id}/campaigns/{campaign_id}", response_model=CampaignOut)
async def get_campaign(
    company_id: str,
    campaign_id: str,
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    row = await db.scalar(select(Campaign).where(Campaign.company_id == company_id, Campaign.id == campaign_id))
    if row is None:
        raise ApiProblem(404, "not_found", "Không tìm thấy chiến dịch.")
    return await _campaign_out(db, row)


@router.get("/workspaces/{company_id}/posts", response_model=PaginatedPosts)
async def list_posts(
    company_id: str,
    campaign_id: str | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=100, ge=1, le=200),
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    statement = select(CampaignPost).where(CampaignPost.company_id == company_id)
    if campaign_id:
        statement = statement.where(CampaignPost.campaign_id == campaign_id)
    total = int(await db.scalar(select(func.count()).select_from(statement.subquery())) or 0)
    rows = (await db.scalars(statement.order_by(CampaignPost.created_at.desc()).offset((page - 1) * page_size).limit(page_size))).all()
    return PaginatedPosts(items=[await _post_out(db, row) for row in rows], total=total, page=page, page_size=page_size)


@router.post(
    "/workspaces/{company_id}/campaigns/{campaign_id}/posts",
    response_model=PostOut,
    status_code=201,
    dependencies=[Depends(require_csrf)],
)
async def create_manual_post(
    company_id: str,
    campaign_id: str,
    request: CreateManualPostRequest,
    user: User = Depends(current_user),
    membership: Membership = Depends(require_permission("post:edit")),
    db: AsyncSession = Depends(get_db),
):
    campaign = await db.scalar(select(Campaign).where(Campaign.company_id == company_id, Campaign.id == campaign_id))
    if campaign is None:
        raise ApiProblem(404, "not_found", "Không tìm thấy chiến dịch.")
    post_id = new_id()
    now = utcnow()
    content = {
        "caption": request.caption,
        "hashtags": request.hashtags,
        "media": [],
        "review": None,
        "citations": [],
    }
    post = CampaignPost(
        id=post_id,
        company_id=company_id,
        campaign_id=campaign_id,
        channel=(campaign.channels_json or ["facebook_page"])[0],
        pillar=request.pillar,
        format=request.format,
        status="draft",
        current_version=1,
        current_json=content,
        created_at=now,
        updated_at=now,
    )
    db.add(post)
    db.add(PostVersion(
        company_id=company_id,
        campaign_id=campaign_id,
        post_id=post_id,
        version=1,
        content_json=content,
        source="human",
        created_by=user.id,
        created_by_name=user.full_name,
        created_at=now,
    ))
    db.add(AuditEvent(company_id=company_id, actor_user_id=user.id, action="post.create", entity_type="post", entity_id=post_id, metadata_json={"campaign_id": campaign_id, "version": 1, "source": "human"}))
    await db.commit()
    return await _post_out(db, post)


@router.get("/workspaces/{company_id}/posts/{post_id}", response_model=PostOut)
async def get_post(
    company_id: str,
    post_id: str,
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    return await _post_out(db, await _get_post(db, company_id, post_id))


@router.get("/workspaces/{company_id}/posts/{post_id}/versions", response_model=PostVersionListOut)
async def list_post_versions(
    company_id: str,
    post_id: str,
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    post = await _get_post(db, company_id, post_id)
    rows = (await db.scalars(select(PostVersion).where(
        PostVersion.company_id == company_id,
        PostVersion.post_id == post.id,
    ).order_by(PostVersion.version.desc()))).all()
    versions = []
    for version in rows:
        content = version.content_json
        approved = await db.scalar(select(PostApproval).where(
            PostApproval.company_id == company_id,
            PostApproval.post_id == post.id,
            PostApproval.version == version.version,
            PostApproval.decision == "approved",
        ).order_by(PostApproval.decided_at.desc()))
        versions.append(PostVersionOut(
            version=version.version,
            caption=str(content.get("caption", "")),
            hashtags=list(content.get("hashtags", [])),
            media=list(content.get("media", [])),
            source=version.source,
            created_by=version.created_by,
            created_by_name=version.created_by_name,
            created_at=version.created_at,
            note=version.note,
            review=content.get("review"),
            approved_at=approved.decided_at if approved else None,
            approved_by=approved.decided_by if approved else None,
        ))
    return PostVersionListOut(
        post_id=post.id,
        versions=versions,
        current_version=post.current_version,
        pending_approval_version=post.pending_approval_version,
    )


@router.patch(
    "/workspaces/{company_id}/posts/{post_id}",
    response_model=PostOut,
    dependencies=[Depends(require_csrf)],
)
async def update_post(
    company_id: str,
    post_id: str,
    request: UpdatePostRequest,
    user: User = Depends(current_user),
    membership: Membership = Depends(require_permission("post:edit")),
    db: AsyncSession = Depends(get_db),
):
    post = await _get_post(db, company_id, post_id, lock=True)
    _check_post_version(post, request.version)
    old = await db.scalar(select(PostVersion).where(PostVersion.company_id == company_id, PostVersion.post_id == post.id, PostVersion.version == post.current_version))
    if old is None:
        raise ApiProblem(500, "content_version_missing", "Không tìm thấy phiên bản nội dung hiện tại.")
    content = dict(old.content_json)
    if request.caption is not None:
        content["caption"] = request.caption
    if request.hashtags is not None:
        content["hashtags"] = request.hashtags
    was_approved = bool(await db.scalar(select(PostApproval.id).where(
        PostApproval.company_id == company_id,
        PostApproval.post_id == post.id,
        PostApproval.decision == "approved",
    ).limit(1)))
    now = utcnow()
    version_number = post.current_version + 1
    post.current_version = version_number
    post.current_json = content
    post.status = "draft"
    post.pending_approval_version = None
    post.requires_reapproval = was_approved
    post.rejection_reason = None
    post.updated_at = now
    db.add(PostVersion(
        company_id=company_id,
        campaign_id=post.campaign_id,
        post_id=post.id,
        version=version_number,
        content_json=content,
        source="human",
        created_by=user.id,
        created_by_name=user.full_name,
        note=request.note,
        created_at=now,
    ))
    db.add(AuditEvent(company_id=company_id, actor_user_id=user.id, action="post.version.create", entity_type="post", entity_id=post.id, metadata_json={"version": version_number, "source": "human"}))
    await db.commit()
    return await _post_out(db, post)


@router.post(
    "/workspaces/{company_id}/posts/{post_id}/submit-approval",
    response_model=PostOut,
    dependencies=[Depends(require_csrf)],
)
async def submit_post_for_approval(
    company_id: str,
    post_id: str,
    request: SubmitApprovalRequest,
    user: User = Depends(current_user),
    membership: Membership = Depends(require_permission("post:edit")),
    db: AsyncSession = Depends(get_db),
):
    post = await _get_post(db, company_id, post_id, lock=True)
    _check_post_version(post, request.version)
    if post.status not in {"draft", "rejected"}:
        raise ApiProblem(409, "state_conflict", "Chỉ bản nháp hoặc bản bị từ chối mới được gửi duyệt.")
    post.status = "needs_review"
    post.pending_approval_version = request.version
    post.rejection_reason = None
    post.updated_at = utcnow()
    db.add(AuditEvent(company_id=company_id, actor_user_id=user.id, action="post.approval.submit", entity_type="post", entity_id=post.id, metadata_json={"version": request.version}))
    await db.commit()
    return await _post_out(db, post)


@router.post(
    "/workspaces/{company_id}/posts/{post_id}/approval",
    response_model=PostOut,
    dependencies=[Depends(require_csrf)],
)
async def decide_post_approval(
    company_id: str,
    post_id: str,
    request: ApprovalRequest,
    user: User = Depends(current_user),
    membership: Membership = Depends(require_permission("post:approve")),
    db: AsyncSession = Depends(get_db),
):
    post = await _get_post(db, company_id, post_id, lock=True)
    _check_post_version(post, request.version)
    if post.status != "needs_review" or post.pending_approval_version != request.version:
        raise ApiProblem(409, "state_conflict", "Phiên bản này không còn nằm trong hàng chờ duyệt.")
    if request.decision == "rejected" and not request.reason:
        raise ApiProblem(422, "validation_error", "Cần ghi lý do khi từ chối bài viết.")
    now = utcnow()
    db.add(PostApproval(
        company_id=company_id,
        post_id=post.id,
        version=request.version,
        decision=request.decision,
        reason=request.reason,
        decided_by=user.id,
        decided_by_name=user.full_name,
        decided_at=now,
    ))
    post.status = request.decision
    post.pending_approval_version = None
    post.requires_reapproval = False
    post.rejection_reason = request.reason if request.decision == "rejected" else None
    post.updated_at = now
    db.add(AuditEvent(company_id=company_id, actor_user_id=user.id, action=f"post.approval.{request.decision}", entity_type="post", entity_id=post.id, metadata_json={"version": request.version}))
    await db.commit()
    return await _post_out(db, post)


@router.get("/workspaces/{company_id}/posts/{post_id}/approvals", response_model=list[ApprovalRecordOut])
async def post_approval_history(
    company_id: str,
    post_id: str,
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    await _get_post(db, company_id, post_id)
    rows = (await db.scalars(select(PostApproval).where(PostApproval.company_id == company_id, PostApproval.post_id == post_id).order_by(PostApproval.decided_at.desc()))).all()
    return [ApprovalRecordOut(id=row.id, post_id=row.post_id, version=row.version, decision=row.decision, reason=row.reason, decided_by=row.decided_by, decided_by_name=row.decided_by_name, decided_at=row.decided_at) for row in rows]


@router.get("/workspaces/{company_id}/approvals", response_model=PaginatedPosts)
async def approval_queue(
    company_id: str,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=100),
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    statement = select(CampaignPost).where(CampaignPost.company_id == company_id, CampaignPost.status == "needs_review")
    total = int(await db.scalar(select(func.count()).select_from(statement.subquery())) or 0)
    rows = (await db.scalars(statement.order_by(CampaignPost.updated_at.asc()).offset((page - 1) * page_size).limit(page_size))).all()
    return PaginatedPosts(items=[await _post_out(db, row) for row in rows], total=total, page=page, page_size=page_size)


@router.post(
    "/workspaces/{company_id}/exports",
    status_code=202,
    dependencies=[Depends(require_csrf)],
)
async def create_export(
    company_id: str,
    request: CreateExportRequest,
    idempotency_key: str | None = Header(default=None, alias="Idempotency-Key", min_length=8, max_length=200),
    user: User = Depends(current_user),
    membership: Membership = Depends(require_permission("export:create")),
    db: AsyncSession = Depends(get_db),
):
    if not idempotency_key:
        raise ApiProblem(400, "idempotency_key_required", "Thiếu Idempotency-Key cho yêu cầu xuất tệp.")
    campaign = await db.scalar(select(Campaign).where(Campaign.company_id == company_id, Campaign.id == request.campaign_id))
    if campaign is None:
        raise ApiProblem(404, "not_found", "Không tìm thấy chiến dịch.")
    existing = await db.scalar(select(Job).where(Job.company_id == company_id, Job.idempotency_key == idempotency_key))
    request_json = request.model_dump(mode="json")
    fingerprint = hashlib.sha256(json.dumps(request_json, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    if existing:
        if (existing.result or {}).get("request_fingerprint") != fingerprint:
            raise ApiProblem(409, "idempotency_conflict", "Idempotency-Key này đã được dùng cho yêu cầu khác.")
        return await accepted_response(db, existing)
    columns = request.columns or list(EXPORT_COLUMNS)
    if not columns or len(set(columns)) != len(columns) or any(column not in EXPORT_COLUMNS for column in columns):
        raise ApiProblem(422, "validation_error", "Danh sách cột xuất không hợp lệ.")
    statement = select(CampaignPost).where(CampaignPost.company_id == company_id, CampaignPost.campaign_id == campaign.id)
    if request.post_ids is not None:
        ids = list(dict.fromkeys(request.post_ids))
        if not ids:
            raise ApiProblem(422, "validation_error", "Cần chọn ít nhất một bài để xuất.")
        statement = statement.where(CampaignPost.id.in_(ids))
    posts = (await db.scalars(statement.order_by(CampaignPost.created_at))).all()
    if request.post_ids is not None and len(posts) != len(set(request.post_ids)):
        raise ApiProblem(404, "not_found", "Một hoặc nhiều bài không thuộc chiến dịch này.")
    if not posts:
        raise ApiProblem(422, "empty_export", "Chiến dịch chưa có bài viết để xuất.")
    rows: list[dict[str, Any]] = []
    versions: dict[str, int] = {}
    for post in posts:
        version = await db.scalar(select(PostVersion).where(
            PostVersion.company_id == company_id,
            PostVersion.post_id == post.id,
            PostVersion.version == post.current_version,
        ))
        if version is None:
            raise ApiProblem(500, "content_version_missing", "Không tìm thấy phiên bản nội dung cần xuất.")
        content = version.content_json
        citations = content.get("citations", [])
        rows.append({
            "post_id": post.id,
            "version": post.current_version,
            "status": post.status,
            "channel": post.channel,
            "pillar": post.pillar,
            "format": post.format,
            "caption": content.get("caption", ""),
            "hashtags": " ".join(content.get("hashtags", [])),
            "scheduled_at": post.scheduled_at.isoformat() if post.scheduled_at else "",
            "citations": " | ".join(f"{item.get('source_id', '')} {item.get('locator', '')}" for item in citations),
        })
        versions[post.id] = post.current_version
    artifact_id = new_id()
    safe_name = "".join(character for character in campaign.name if character.isalnum() or character in " -_").strip()[:100] or "campaign"
    extension = request.format
    filename = f"{safe_name}-{artifact_id[:8]}.{extension}"
    content_bytes = _export_bytes(rows, columns, extension)
    object_key = f"exports/{company_id}/{artifact_id}.{extension}"
    try:
        await storage.put(object_key, content_bytes)
    except Exception:
        raise ApiProblem(503, "export_storage_unavailable", "Không lưu được tệp xuất. Vui lòng thử lại.", retryable=True)
    now = utcnow()
    artifact = ExportArtifact(
        id=artifact_id,
        company_id=company_id,
        campaign_id=campaign.id,
        created_by=user.id,
        format=extension,
        filename=filename,
        file_size_bytes=len(content_bytes),
        object_key=object_key,
        content_sha256=hashlib.sha256(content_bytes).hexdigest(),
        post_ids_json=[post.id for post in posts],
        version_snapshot_json=versions,
        created_at=now,
        updated_at=now,
    )
    job = Job(
        company_id=company_id,
        created_by=user.id,
        kind="export",
        title=f"Xuất {len(posts)} bài: {campaign.name}",
        status="succeeded",
        progress=100,
        result={"export_id": artifact_id, "filename": filename, "size": len(content_bytes), "request_fingerprint": fingerprint},
        idempotency_key=idempotency_key,
        started_at=now,
        finished_at=now,
        created_at=now,
        updated_at=now,
    )
    db.add(artifact)
    db.add(job)
    db.add(AuditEvent(company_id=company_id, actor_user_id=user.id, action="campaign.export", entity_type="campaign", entity_id=campaign.id, metadata_json={"export_id": artifact_id, "post_count": len(posts), "format": extension, "sha256": artifact.content_sha256}))
    await db.commit()
    return await accepted_response(db, job)


@router.get("/workspaces/{company_id}/exports/{export_id}", response_model=ExportOut)
async def get_export(
    company_id: str,
    export_id: str,
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    artifact = await db.scalar(select(ExportArtifact).where(ExportArtifact.company_id == company_id, ExportArtifact.id == export_id))
    if artifact is None:
        raise ApiProblem(404, "not_found", "Không tìm thấy tệp xuất.")
    return ExportOut(id=artifact.id, campaign_id=artifact.campaign_id, format=artifact.format, status="ready", download_url=f"/workspaces/{company_id}/exports/{artifact.id}/download", filename=artifact.filename, size=artifact.file_size_bytes, created_at=artifact.created_at)


@router.get("/workspaces/{company_id}/exports/{export_id}/download")
async def download_export(
    company_id: str,
    export_id: str,
    user: User = Depends(current_user),
    db: AsyncSession = Depends(get_db),
):
    await membership_for(company_id, user, db)
    artifact = await db.scalar(select(ExportArtifact).where(ExportArtifact.company_id == company_id, ExportArtifact.id == export_id))
    if artifact is None:
        raise ApiProblem(404, "not_found", "Không tìm thấy tệp xuất.")
    try:
        body = await storage.read(artifact.object_key) if isinstance(storage, S3ObjectStorage) else storage.path(artifact.object_key).read_bytes()
    except Exception:
        raise ApiProblem(404, "export_file_missing", "Tệp xuất hiện không còn trong kho lưu trữ.")
    media_type = "text/csv; charset=utf-8" if artifact.format == "csv" else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return Response(
        content=body,
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(artifact.filename)}", "Cache-Control": "private, no-store", "X-Content-SHA256": artifact.content_sha256},
    )


@router.post(
    "/workspaces/{company_id}/posts/generate",
    response_model=GenerateContentResponse,
    status_code=503,
    dependencies=[Depends(require_csrf)],
)
async def generate_content(
    company_id: str,
    request: GenerateContentRequest,
    user: User = Depends(current_user),
    membership: Membership = Depends(require_permission("post:generate")),
    db: AsyncSession = Depends(get_db),
):
    # This endpoint stays fail-closed until the workspace owner approves
    # sending brand-profile and retrieved source text to DeepSeek.
    raise ApiProblem(
        503,
        "provider_approval_required",
        "Sinh nội dung đang tạm dừng. Cần xác nhận rằng Brand Profile và nguồn tài liệu sẽ được gửi tới DeepSeek để xử lý.",
        retryable=False,
    )
